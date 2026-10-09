"""ATS-safe resume export: single column, standard headings, no tables/images/text boxes.

The renderers take a TailoredResume (ai/schemas.py) or any object with the same attributes.
All field access goes through `_sections`.
"""

import io
import re
import unicodedata
from pathlib import Path
from xml.sax.saxutils import escape

from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from jobfinder.letters.docx_export import new_document, to_bytes

# The built-in Helvetica is Latin-1 only, so names like "Łukasz" or Cyrillic/Greek text would render
# as blanks. DejaVu Sans (bundled, see fonts/LICENSE-DejaVu.txt) covers Latin Extended, Vietnamese,
# Cyrillic and Greek. It does NOT cover CJK: those characters render as the font's missing-glyph box
# (no crash) and the DOCX export should be preferred for them.
FONT_DIR = Path(__file__).parent / "fonts"
FONT, FONT_BOLD = "DejaVuSans", "DejaVuSans-Bold"


def _register_fonts() -> None:
    registered = pdfmetrics.getRegisteredFontNames()
    for name, file in ((FONT, "DejaVuSans.ttf"), (FONT_BOLD, "DejaVuSans-Bold.ttf")):
        if name not in registered:
            pdfmetrics.registerFont(TTFont(name, FONT_DIR / file))


_NON_BMP = re.compile("[^\U00000000-\U0000ffff]")
# XML 1.0 forbids C0 controls other than \t \n \r, lone surrogates and U+FFFE/U+FFFF; python-docx
# raises on them. \x0b/\x0c are soft line/page breaks in pasted text, so they become spaces.
_XML_INVALID = re.compile("[\x00-\x08\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_BREAKS = re.compile("[\x0b\x0c]")


def _clean(text: str) -> str:
    return _XML_INVALID.sub("", _BREAKS.sub(" ", text))


def _pdf_text(text: str) -> str:
    """NFC-compose (reportlab does no mark positioning) and replace non-BMP characters, which
    reportlab's TTF subsetting would otherwise truncate into different, wrong characters."""
    return escape(_NON_BMP.sub("\ufffd", unicodedata.normalize("NFC", text)))


HEADINGS = ("Summary", "Skills", "Experience", "Projects", "Education")


def _join(parts, sep=" | ") -> str:
    return sep.join(str(p).strip() for p in parts if p and str(p).strip())


def _bullets(summary, bullets, technologies) -> list[str]:
    out = ([summary] if summary else []) + list(bullets or [])
    if technologies:
        out.append("Technologies: " + ", ".join(technologies))
    return out


def _sections(resume) -> tuple[str, list[str], list[tuple[str, list[tuple[str, list[str]]]]]]:
    """Normalise to (name, contact lines, [(heading, [(entry line, bullets)])]); empty sections are dropped."""
    c = resume.contact
    contact = [_join([c.email, c.phone, c.location]), _join(c.links or [])]
    sections: list[tuple[str, list[tuple[str, list[str]]]]] = []
    if resume.summary:
        sections.append(("Summary", [(resume.summary, [])]))
    if resume.skills:
        sections.append(("Skills", [(", ".join(resume.skills), [])]))
    if resume.experience:
        sections.append(
            (
                "Experience",
                [
                    (
                        _join([e.title, e.company, e.location, _join([e.start, e.end], " - ")]),
                        _bullets(e.summary, e.bullets, e.technologies),
                    )
                    for e in resume.experience
                ],
            )
        )
    if resume.projects:
        sections.append(
            (
                "Projects",
                [
                    (_join([p.name, p.url]), _bullets(p.description, [], p.technologies))
                    for p in resume.projects
                ],
            )
        )
    if resume.education:
        sections.append(
            (
                "Education",
                [
                    (
                        _join([_join([e.degree, e.field], ", "), e.school, _join([e.start, e.end], " - ")]),
                        [f"GPA {e.gpa}"] if e.gpa else [],
                    )
                    for e in resume.education
                ],
            )
        )
    return (
        _clean(c.name),
        [y for y in (_clean(x) for x in contact) if y],
        [
            (h, [(_clean(line), [_clean(b) for b in bullets]) for line, bullets in entries])
            for h, entries in sections
        ],
    )


def resume_to_text(resume) -> str:
    name, contact, sections = _sections(resume)
    lines = [name, *contact]
    for heading, entries in sections:
        lines += ["", heading]
        for line, bullets in entries:
            if line:
                lines.append(line)
            lines += [f"- {b}" for b in bullets]
    return "\n".join(lines) + "\n"


def render_docx(resume) -> bytes:
    name, contact, sections = _sections(resume)
    d = new_document()
    title = d.add_paragraph()
    run = title.add_run(name)
    run.bold = True
    run.font.size = Pt(16)
    title.paragraph_format.space_after = Pt(2)
    for line in contact:
        d.add_paragraph(line).paragraph_format.space_after = Pt(0)
    for heading, entries in sections:
        d.add_heading(heading, level=1)
        for line, bullets in entries:
            if line:
                para = d.add_paragraph(line)
                para.paragraph_format.space_after = Pt(2)
                if bullets:
                    para.runs[0].bold = True
            for b in bullets:
                d.add_paragraph(b, style="List Bullet").paragraph_format.space_after = Pt(0)
    # Heading 1 defaults to a colored theme font; keep fonts plain and black for parsers
    st = d.styles["Heading 1"]
    st.font.name = "Calibri"
    st.font.size = Pt(13)
    st.font.bold = True
    st.font.color.rgb = None
    return to_bytes(d)


def render_pdf(resume) -> bytes:
    name, contact, sections = _sections(resume)
    _register_fonts()
    base = ParagraphStyle("base", fontName=FONT, fontSize=10, leading=13, spaceAfter=2)
    title = ParagraphStyle("title", parent=base, fontName=FONT_BOLD, fontSize=16, leading=19)
    head = ParagraphStyle(
        "head", parent=base, fontName=FONT_BOLD, fontSize=12.5, leading=15, spaceBefore=8
    )
    entry = ParagraphStyle("entry", parent=base, fontName=FONT_BOLD)
    story = [Paragraph(_pdf_text(name), title), *(Paragraph(_pdf_text(line), base) for line in contact)]
    for heading, entries in sections:
        story.append(Paragraph(_pdf_text(heading), head))
        for line, bullets in entries:
            if line:
                story.append(Paragraph(_pdf_text(line), entry if bullets else base))
            if bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(_pdf_text(b), base)) for b in bullets],
                        bulletType="bullet",
                        start="-",
                        leftIndent=14,
                    )
                )
        story.append(Spacer(1, 2))
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=name,
    ).build(story)
    return buf.getvalue()
