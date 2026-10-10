import io
import re

import docx
from docx.shared import Inches, Pt

# XML 1.0 forbids C0 controls other than \t \n \r, lone surrogates and U+FFFE/U+FFFF; python-docx
# raises on them. \x0b/\x0c are soft line/page breaks in pasted text, so they become spaces.
_XML_INVALID = re.compile("[\x00-\x08\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_BREAKS = re.compile("[\x0b\x0c]")


def clean_text(text: str) -> str:
    return _XML_INVALID.sub("", _BREAKS.sub(" ", text))


def new_document() -> docx.document.Document:
    """A blank single-column document with Calibri 11 and 1-inch margins."""
    d = docx.Document()
    for section in d.sections:
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return d


def to_bytes(d: docx.document.Document) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_docx(content: str) -> bytes:
    """A plain, ATS-friendly Word letter: Calibri 11, 1-inch margins, blank-line separated paragraphs."""
    d = new_document()
    for block in re.split(r"\n\s*\n", clean_text(content).strip()):
        lines = block.strip().split("\n")
        para = d.add_paragraph()
        para.paragraph_format.space_after = Pt(10)
        for i, line in enumerate(lines):
            run = para.add_run(clean_text(line).strip())
            if i < len(lines) - 1:
                run.add_break()  # keep "Sincerely,\nName" on adjacent lines
    return to_bytes(d)


def filename(company: str, title: str) -> str:
    base = re.sub(r"[^A-Za-z0-9 ._-]+", "", f"Cover Letter - {company} - {title}").strip()
    base = re.sub(r"\s+", " ", base)[:120]
    return f"{base}.docx"
