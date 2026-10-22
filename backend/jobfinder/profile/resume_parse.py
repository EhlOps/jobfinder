"""Extract plain text from uploaded resume/portfolio files."""
import asyncio
import io
import threading
from concurrent.futures import ThreadPoolExecutor
from itertools import islice

from jobfinder.profile.formats import ALLOWED_EXTENSIONS  # noqa: F401  (re-exported)

MAX_CHARS = 60_000
MAX_PDF_PAGES = 30  # a resume is a few pages; anything longer is cut off, not parsed to the end
PARSE_TIMEOUT_S = 20.0
MAX_PARSES = 2
SLOT_POLL_S = 0.25  # how often worker-priority looks for a free slot
_slots = threading.BoundedSemaphore(MAX_PARSES)
_pool = ThreadPoolExecutor(max_workers=MAX_PARSES, thread_name_prefix="parse")


def extract_text(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if ext == ".pdf":
        import pdfplumber

        parts, total = [], 0
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in islice(pdf.pages, MAX_PDF_PAGES):
                parts.append(page.extract_text() or "")
                total += len(parts[-1])
                if total > MAX_CHARS:
                    break
        text = "\n\n".join(parts)
    elif ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return text.strip()[:MAX_CHARS]


async def extract_text_in_worker(data: bytes, ext: str) -> str:
    """extract_text for worker-priority's parse_document (#51) on a dedicated two-thread pool with a time limit,
    waiting for a free slot. A thread can't be killed, so a slot is released only when its parse really
    finishes (not on timeout): slow files make later parses wait instead of piling up more threads, and
    ingestion's DNS lookups and to_thread calls on the default executor never queue behind a stuck parse."""
    while not _slots.acquire(blocking=False):
        await asyncio.sleep(SLOT_POLL_S)
    try:
        fut = asyncio.get_running_loop().run_in_executor(_pool, _run, data, ext)
    except BaseException:
        _slots.release()
        raise
    return await asyncio.wait_for(fut, PARSE_TIMEOUT_S)


def _run(data: bytes, ext: str) -> str:
    try:
        return extract_text(data, ext)
    finally:
        _slots.release()
