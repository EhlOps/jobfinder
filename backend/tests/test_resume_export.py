import io
from types import SimpleNamespace as NS

import docx
import pdfplumber

from jobfinder.resumes.export import HEADINGS, render_docx, render_pdf, resume_to_text


def _resume():
    return NS(
        contact=NS(
            name="Ada Lovelace",
            email="ada@example.com",
            phone="555-0100",
            location="Boston, MA",
            links=["github.com/ada"],
        ),
        summary="Engineer who builds reliable analytical engines.",
        skills=["Python", "SQL", "FastAPI"],
        experience=[
            NS(
                company="Acme",
                role="Software Engineer",
                start="2022",
                end="Present",
                location="Remote",
                bullets=["Built a billing service", "Cut latency by 40% & saved cost"],
            ),
        ],
        projects=[
            NS(name="Engine", bullets=["Open-source <calculator>"], link="github.com/ada/engine")
        ],
        education=[
            NS(
                school="Northeastern",
                degree="BS Computer Science",
                dates="2018-2022",
                details=["GPA 3.9"],
            )
        ],
    )


def _in_order(text: str, needles: list[str]) -> None:
    pos = 0
    for n in needles:
        i = text.find(n, pos)
        assert i >= 0, f"{n!r} missing or out of order"
        pos = i + len(n)


ORDER = [
    "Ada Lovelace",
    "ada@example.com",
    "Summary",
    "reliable analytical engines",
    "Skills",
    "Python, SQL, FastAPI",
    "Experience",
    "Software Engineer",
    "Acme",
    "Built a billing service",
    "Cut latency by 40% & saved cost",
    "Projects",
    "Engine",
    "Open-source <calculator>",
    "Education",
    "BS Computer Science",
    "Northeastern",
    "GPA 3.9",
]


def test_docx_is_ats_safe():
    d = docx.Document(io.BytesIO(render_docx(_resume())))
    assert not d.tables
    assert not d.inline_shapes
    xml = d.element.xml
    assert "txbxContent" not in xml and "<w:drawing" not in xml and "<w:pict" not in xml
    assert [p.text for p in d.paragraphs if p.style.name == "Heading 1"] == list(HEADINGS)
    assert len(d.sections) == 1
    _in_order("\n".join(p.text for p in d.paragraphs), ORDER)


def test_pdf_text_matches_in_order():
    with pdfplumber.open(io.BytesIO(render_pdf(_resume()))) as pdf:
        text = "\n".join(page.extract_text() for page in pdf.pages)
        assert not any(page.images for page in pdf.pages)
    _in_order(text, ORDER)


def test_empty_sections_dropped_and_text_export():
    r = _resume()
    r.projects = []
    r.contact.links = []
    text = resume_to_text(r)
    assert "Projects" not in text
    _in_order(text, [n for n in ORDER if n not in ("Projects", "Engine", "Open-source <calculator>")])
    d = docx.Document(io.BytesIO(render_docx(r)))
    assert "Projects" not in [p.text for p in d.paragraphs]
    assert render_pdf(r).startswith(b"%PDF")
