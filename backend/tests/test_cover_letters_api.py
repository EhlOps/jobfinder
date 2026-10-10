import io

import docx
from sqlalchemy.ext.asyncio import async_sessionmaker

from jobfinder.letters.docx_export import build_docx, filename
from jobfinder.models import Job, JobMatch
from jobfinder.scheduling import queue
from jobfinder.worker import run_task

LETTER = "Dear Acme hiring team,\n\nI built things.\n\nSincerely,\nJane Doe"
N = 0


class FakeAI:
    def __init__(self, text=LETTER):
        self.text, self.calls = text, []

    async def cover_letter(self, **kw):
        self.calls.append(kw)
        return self.text


async def setup(client, engine, *, background=True):
    """Sign up, give the user a profile, and create one scored match. Returns (match_id, job_id)."""
    global N
    N += 1
    uid = (await client.get("/api/auth/me")).json()["id"]
    if background:
        await client.put("/api/profile/background", json={"name": "Jane Doe", "experience": [{"company": "Stripe"}]})
        await client.post("/api/profile/facts", json={"question": "Used Go?", "answer": "Yes"})
    async with async_sessionmaker(engine, expire_on_commit=False)() as db:
        job = Job(source="greenhouse", external_id=f"c{N}", company_name="Acme/Co", title="Backend Engineer (II)", url="https://x",
                  location="Boston", description_text="Build services", salary_min=150000, salary_max=200000, dedupe_hash=f"c{N}")
        db.add(job)
        await db.flush()
        m = JobMatch(user_id=uid, job_id=job.id, profile_version=1, prefilter_score=50, llm_score=80, confidence=0.9,
                     verdict="strong", reasons=["Go at Stripe"], gaps=["k8s"], unknowns=[], status="new")
        db.add(m)
        await db.commit()
        return m.id, job.id


async def draft(client, engine, mid, ai=None, **body):
    r = await client.post(f"/api/matches/{mid}/cover-letter", json=body)
    assert r.status_code == 202, r.text
    await run_task(r.json()["task_id"], ai or FakeAI(), async_sessionmaker(engine, expire_on_commit=False))
    return r.json()["task_id"]


async def test_generate_flow_and_prompt_inputs(authed, engine):
    mid, _ = await setup(authed, engine)
    assert (await authed.get(f"/api/matches/{mid}/cover-letter")).json() == {"letter": None, "pending_task_id": None}

    ai = FakeAI()
    tid = await draft(authed, engine, mid, ai, tone="warm", notes="mention Go")
    t = (await authed.get(f"/api/tasks/{tid}")).json()
    assert t["status"] == "done" and t["result"]["words"] > 3

    call = ai.calls[0]
    assert call["tone"] == "warm" and call["notes"] == "mention Go" and call["name"] == "Jane Doe"
    assert call["analysis"] == {"strengths": ["Go at Stripe"], "gaps": ["k8s"]}
    assert call["facts"] == [("Used Go?", "Yes")] and "salary" not in call["job"] and call["job"]["company"] == "Acme/Co"

    st = (await authed.get(f"/api/matches/{mid}/cover-letter")).json()
    assert st["pending_task_id"] is None
    assert st["letter"]["content"] == LETTER and st["letter"]["tone"] == "warm" and st["letter"]["edited"] is False
    assert st["letter"]["words"] == len(LETTER.split()) and st["letter"]["stale"] is False


async def test_pending_task_is_reported_and_deduped(authed, engine):
    mid, _ = await setup(authed, engine)
    a = await authed.post(f"/api/matches/{mid}/cover-letter", json={"tone": "warm"})
    b = await authed.post(f"/api/matches/{mid}/cover-letter", json={"tone": "warm"})
    assert a.json()["task_id"] == b.json()["task_id"]                         # identical request: same task
    c = await authed.post(f"/api/matches/{mid}/cover-letter", json={"tone": "concise"})
    assert c.json()["task_id"] != a.json()["task_id"]                         # different options: new task
    assert (await authed.get(f"/api/matches/{mid}/cover-letter")).json()["pending_task_id"] == c.json()["task_id"]


async def test_edit_then_redraft_needs_force(authed, engine):
    mid, _ = await setup(authed, engine)
    await draft(authed, engine, mid)
    url = f"/api/matches/{mid}/cover-letter"

    same = await authed.put(url, json={"content": LETTER})
    assert same.json()["edited"] is False                                    # unchanged text isn't an edit
    r = await authed.put(url, json={"content": "Dear team,\r\n\r\nMy own words.\r\n\r\nThanks,\r\nJane  "})
    assert r.status_code == 200 and r.json()["edited"] is True
    assert r.json()["content"] == "Dear team,\n\nMy own words.\n\nThanks,\nJane"

    assert (await authed.post(url, json={})).status_code == 409                # would overwrite edits
    await draft(authed, engine, mid, FakeAI("Dear Acme,\n\nRedrafted."), force=True)
    st = (await authed.get(url)).json()["letter"]
    assert st["content"] == "Dear Acme,\n\nRedrafted." and st["edited"] is False

    assert (await authed.post(url, json={})).status_code == 202               # not edited: no force needed
    assert (await authed.put(url, json={"content": ""})).status_code == 422
    assert (await authed.put(url, json={"content": "x" * 10_001})).status_code == 422


async def test_validation_and_prerequisites(authed, engine):
    mid, _ = await setup(authed, engine)
    url = f"/api/matches/{mid}/cover-letter"
    assert (await authed.post(url, json={"tone": "shouty"})).status_code == 422
    assert (await authed.post(url, json={"notes": "x" * 501})).status_code == 422
    assert (await authed.put(url, json={"content": "hi"})).status_code == 404        # nothing to edit yet
    assert (await authed.get(url + "/download")).status_code == 404


async def test_requires_profile(authed, engine):
    mid, _ = await setup(authed, engine, background=False)
    assert (await authed.post(f"/api/matches/{mid}/cover-letter", json={})).status_code == 422


async def test_download_docx(authed, engine):
    mid, _ = await setup(authed, engine)
    await draft(authed, engine, mid)
    r = await authed.get(f"/api/matches/{mid}/cover-letter/download")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Cover%20Letter%20-%20AcmeCo%20-%20Backend%20Engineer%20II.docx" in r.headers["content-disposition"]
    paragraphs = [p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs]
    assert paragraphs[0] == "Dear Acme hiring team," and paragraphs[1] == "I built things." and "Jane Doe" in paragraphs[2]


def test_docx_keeps_signature_lines_together_and_handles_odd_text():
    d = docx.Document(io.BytesIO(build_docx("Hi\n\n\n\nBody line one\nline two\n\nSincerely,\nJane")))
    texts = [p.text for p in d.paragraphs]
    assert texts == ["Hi", "Body line one\nline two", "Sincerely,\nJane"]
    assert filename("A/B Inc.", "Engineer, Backend (Remote)") == "Cover Letter - AB Inc. - Engineer Backend Remote.docx"
    assert len(filename("x" * 300, "y")) <= 125


def test_docx_survives_control_characters():
    out = build_docx("Dear \x00team,\x0bthanks\n\nSincerely,\ufffeJane\ud800")
    d = docx.Document(io.BytesIO(out))
    assert [p.text for p in d.paragraphs] == ["Dear team, thanks", "Sincerely,Jane"]


async def test_delete_and_flags_in_list(authed, engine):
    mid, _ = await setup(authed, engine)
    assert (await authed.get("/api/matches")).json()["items"][0]["has_cover_letter"] is False
    await draft(authed, engine, mid)
    assert (await authed.get("/api/matches")).json()["items"][0]["has_cover_letter"] is True
    assert (await authed.get(f"/api/matches/{mid}")).json()["has_cover_letter"] is True
    assert (await authed.delete(f"/api/matches/{mid}/cover-letter")).status_code == 204
    assert (await authed.get(f"/api/matches/{mid}/cover-letter")).json()["letter"] is None
    assert (await authed.get(f"/api/matches/{mid}")).json()["has_cover_letter"] is False


async def test_letters_are_private_to_their_owner(authed, client, engine, make_account):
    mid, _ = await setup(authed, engine)
    await draft(authed, engine, mid)
    await client.post("/api/auth/logout")
    await make_account("other@x.com", "password-1234")
    base = f"/api/matches/{mid}/cover-letter"
    assert (await client.get(base)).status_code == 404
    assert (await client.post(base, json={})).status_code == 404
    assert (await client.put(base, json={"content": "hijack"})).status_code == 404
    assert (await client.delete(base)).status_code == 404
    assert (await client.get(base + "/download")).status_code == 404


async def test_requires_login(client):
    assert (await client.get("/api/matches/1/cover-letter")).status_code == 401


async def test_stale_flag_when_profile_changes_after_draft(authed, engine):
    mid, _ = await setup(authed, engine)
    await draft(authed, engine, mid)
    await authed.put("/api/profile/status", json={"target_roles": ["x"]})        # bumps the profile version
    assert (await authed.get(f"/api/matches/{mid}/cover-letter")).json()["letter"]["stale"] is True


async def test_queue_dedupe_considers_payload(engine):
    async with async_sessionmaker(engine, expire_on_commit=False)() as db:
        a = await queue.enqueue(db, "k", {"m": 1}, user_id=None, dedupe=True)
        assert (await queue.enqueue(db, "k", {"m": 1}, user_id=None, dedupe=True)).id == a.id
        assert (await queue.enqueue(db, "k", {"m": 2}, user_id=None, dedupe=True)).id != a.id
        assert (await queue.enqueue(db, "k", None, user_id=None, dedupe=True)).id not in (a.id,)
